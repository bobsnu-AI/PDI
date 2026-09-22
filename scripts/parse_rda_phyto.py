"""Parse the three RDA Korean phytochemical DBs into long format (food_kr, compound, value)."""
import csv, os, re, sys
import openpyxl

DB = sys.argv[1]
MAP = os.path.join(DB, "mapping")
os.makedirs(MAP, exist_ok=True)
R = os.path.join(DB, "kr_rda")


def clean(x):
    return re.sub(r"\s+", " ", str(x)).strip() if x is not None else ""


def num(v):
    """'-' and '' mean not detected / no data; return float or None."""
    s = clean(v)
    if s in ("", "-", "nd", "ND", "N/A", "tr"):
        return None
    s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def parse_wide(path, dbname, skip_sheets=("DB 설명",)):
    """Sheets shaped: row0 = ['성분명', food1, food2, ...]; rows = compound x value."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        if ws.title in skip_sheets:
            continue
        data = list(ws.iter_rows(values_only=True))
        if not data:
            continue
        hdr = data[0]
        foods = {i: clean(hdr[i]) for i in range(1, len(hdr)) if clean(hdr[i])}
        if not foods:
            continue
        for r in data[1:]:
            cmp_ = clean(r[0])
            if not cmp_ or cmp_.startswith("*") or "총 " in cmp_ or cmp_.startswith("="):
                continue
            for i, food in foods.items():
                if i >= len(r):
                    continue
                v = num(r[i])
                if v is None:
                    continue
                rows.append({"db": dbname, "sheet": ws.title, "food_kr": food,
                             "compound": cmp_, "value": v, "unit": "mg/100g DW"})
    wb.close()
    return rows


def parse_saponin(path):
    """Per-plant sheets: a '식품명' row names the foods per column; compound rows follow."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        if ws.title in ("DB 설명", "수록 식품 정보"):
            continue
        data = list(ws.iter_rows(values_only=True))
        namerow = next((r for r in data if clean(r[3] if len(r) > 3 else "") == "식품명"), None)
        if namerow is None:
            continue
        foods = {i: clean(namerow[i]) for i in range(4, len(namerow)) if clean(namerow[i])}
        start = data.index(namerow)
        for r in data[start + 1:]:
            cmp_ = clean(r[2] if len(r) > 2 else "") or clean(r[1] if len(r) > 1 else "")
            if not cmp_ or cmp_ in ("Food Description", "식품명", "학명") or cmp_.startswith("*"):
                continue
            if "총 " in cmp_ or cmp_.startswith("="):
                continue
            for i, food in foods.items():
                if i >= len(r):
                    continue
                v = num(r[i])
                if v is None:
                    continue
                rows.append({"db": "rda_saponin", "sheet": ws.title, "food_kr": food,
                             "compound": cmp_, "value": v, "unit": "mg/100g DW"})
    wb.close()
    return rows


all_rows = []
all_rows += parse_wide(os.path.join(R, "플라보노이드 Data Base 1.0.xlsx"), "rda_flavonoid")
all_rows += parse_wide(os.path.join(R, "공공데이터_페놀산 DB.xlsx"), "rda_phenolic_acid")
all_rows += parse_saponin(os.path.join(R, "RDA 기능성분 DB 3.0 사포닌 편.xlsx"))

with open(os.path.join(MAP, "rda_phyto_long.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, ["db", "sheet", "food_kr", "compound", "value", "unit"])
    w.writeheader()
    w.writerows(all_rows)

from collections import Counter, defaultdict
print(f"total rows: {len(all_rows)}")
for db in ("rda_flavonoid", "rda_phenolic_acid", "rda_saponin"):
    rs = [r for r in all_rows if r["db"] == db]
    pos = [r for r in rs if r["value"] > 0]
    print(f"  {db:<20} rows={len(rs):>6} (>0: {len(pos):>6})  foods={len({r['food_kr'] for r in rs}):>4} "
          f"(>0: {len({r['food_kr'] for r in pos}):>4})  compounds={len({r['compound'] for r in rs}):>4} "
          f"(>0: {len({r['compound'] for r in pos}):>4})")
    print(f"      sheets: {sorted({r['sheet'] for r in rs})}")
print("\nsample foods per db:")
for db in ("rda_flavonoid", "rda_phenolic_acid", "rda_saponin"):
    fs = sorted({r["food_kr"] for r in all_rows if r["db"] == db and r["value"] > 0})
    print(f"  {db}: {fs[:14]}")
