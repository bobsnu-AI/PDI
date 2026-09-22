"""Step A/B: build the Korean base food list (RDA 10.4 plant foods) and each DB's food roster."""
import csv, glob, json, os, re, sys
import openpyxl

csv.field_size_limit(10 ** 8)
DB = sys.argv[1]
MAP = os.path.join(DB, "mapping")
os.makedirs(MAP, exist_ok=True)

PLANT_GROUPS = {          # user-specified plant groups; 향신료 -> 조미료류, 음료 -> 음료류+차류
    "곡류 및 그 제품": "곡류",
    "두류": "두류",
    "견과류 및 종실류": "견과종실",
    "채소류": "채소",
    "과일류": "과일",
    "버섯류": "버섯",
    "해조류": "해조",
    "음료류": "음료",
    "차류": "음료(차)",
    "조미료류": "향신료·조미료",
}

# ---------- RDA 10.4 + 부록2 ----------
wb = openpyxl.load_workbook(os.path.join(DB, "kr_rda/식품성분표(10개정판).xlsx"), read_only=True)
ws = wb["국가표준식품성분 Database 10.4"]
it = ws.iter_rows(values_only=True)
next(it); next(it); next(it)                      # 3 header rows
rda = []
for r in it:
    if not r or not r[2]:
        continue
    rda.append({"idx": str(r[0] or "").strip(), "group": str(r[2]).strip(),
                "kr": str(r[3] or "").strip(), "src": str(r[4] or "").strip(),
                "kcal": r[5], "beta_carotene_ug": r[35]})

app2 = None
for s in wb.sheetnames:
    if s.startswith("부록2"):
        app2 = wb[s]
en = {}
for i, r in enumerate(app2.iter_rows(values_only=True)):
    if i == 0:
        continue
    if not r[0]:
        continue
    en[str(r[0]).strip()] = {"code": str(r[1] or "").strip(), "kr": str(r[2] or "").strip(),
                             "en": str(r[3] or "").strip(), "sci": str(r[4] or "").strip()}
wb.close()

base = []
for f in rda:
    if f["group"] not in PLANT_GROUPS:
        continue
    e = en.get(f["idx"], {})
    base.append({
        "rda_idx": f["idx"], "rda_code": e.get("code", ""), "group": f["group"],
        "group_short": PLANT_GROUPS[f["group"]], "name_kr": f["kr"],
        "name_en": e.get("en", ""), "sci_name": e.get("sci", ""),
        "kcal": f["kcal"], "beta_carotene_ug": f["beta_carotene_ug"], "source": f["src"],
    })
with open(os.path.join(MAP, "base_foods_rda104.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, list(base[0]))
    w.writeheader()
    w.writerows(base)
print(f"base plant foods (RDA 10.4): {len(base)} / all RDA items {len(rda)}")
from collections import Counter
for g, n in Counter(b["group"] for b in base).most_common():
    print(f"   {g:<16} {n}")
print(f"   with 영문명: {sum(1 for b in base if b['name_en'])}, with 학명: {sum(1 for b in base if b['sci_name'])}")
print(f"   distinct 학명: {len({b['sci_name'] for b in base if b['sci_name']})}")

# ---------- KNHANES food codes (intake-frequency proxy / RDA code bridge) ----------
kn = []
for p in sorted(glob.glob(os.path.join(DB, "kr_knhanes/*코드자료집*.xlsx"))):
    year = re.search(r"\((\d{4})\)", p)
    w2 = openpyxl.load_workbook(p, read_only=True)
    if "식품코드" not in w2.sheetnames:
        w2.close()
        continue
    s = w2["식품코드"]
    for i, r in enumerate(s.iter_rows(values_only=True)):
        if i == 0 or not r[0]:
            continue
        kn.append({"year": year.group(1) if year else "", "knhanes_code": str(r[0]).strip(),
                   "name_kr": str(r[1] or "").strip(), "rda_code": str(r[2] or "").strip()})
    w2.close()
latest = max((k["year"] for k in kn), default="")
kn_latest = [k for k in kn if k["year"] == latest]
with open(os.path.join(MAP, "knhanes_food_codes.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, ["year", "knhanes_code", "name_kr", "rda_code"])
    w.writeheader()
    w.writerows(kn)
print(f"\nKNHANES food codes: {len(kn)} rows, latest={latest} ({len(kn_latest)} foods), "
      f"with 농진청코드 {sum(1 for k in kn_latest if k['rda_code'])}")

# ---------- foreign DB food rosters ----------
rosters = {}

pe = list(csv.DictReader(open(os.path.join(DB, "phenol_explorer/foods.csv"), encoding="utf-8", errors="replace")))
rosters["phenol_explorer"] = [{"id": r["id"], "name_en": r["name"], "sci_name": r["food_source_scientific_name"],
                               "group": r["food_group"]} for r in pe]

for key, sub in (("usda_flavonoid", "flavonoid_csv"), ("usda_isoflavone", "isoflavone_csv"),
                 ("usda_proanthocyanidin", "proanthocyanidin_csv")):
    fd = list(csv.DictReader(open(os.path.join(DB, "usda", sub, "FOOD_DES.csv"), encoding="utf-8", errors="replace")))
    k = "Long_Desc" if "Long_Desc" in fd[0] else next(c for c in fd[0] if "Desc" in c)
    sci = "SciName" if "SciName" in fd[0] else ""
    rosters[key] = [{"id": (r.get("NDB_No") or r.get("NDB No") or ""), "name_en": r[k], "sci_name": (r.get(sci, "") if sci else ""),
                     "group": r.get("FdGrp_Cd", "")} for r in fd]

car = list(csv.DictReader(open(os.path.join(DB, "usda/carotenoids.csv"), encoding="utf-8", errors="replace")))
rosters["usda_sr_carotenoid"] = [{"id": r["ndb_number"], "name_en": r["description"], "sci_name": "",
                                  "group": r["food_category"]} for r in car]

FB = os.path.join(DB, "foodb/foodb_2020_04_07_csv")
fb = list(csv.DictReader(open(os.path.join(FB, "Food.csv"), encoding="utf-8", errors="replace")))
rosters["foodb"] = [{"id": r["id"], "name_en": r["name"], "sci_name": r["name_scientific"],
                     "group": r["food_group"]} for r in fb]

with open(os.path.join(MAP, "db_food_rosters.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(["db", "food_id", "name_en", "sci_name", "group"])
    for db, rows in rosters.items():
        for r in rows:
            w.writerow([db, r["id"], r["name_en"], r["sci_name"], r["group"]])
print("\nDB food rosters:")
for db, rows in rosters.items():
    print(f"   {db:<24} {len(rows):>5} foods, with sci_name {sum(1 for r in rows if (r['sci_name'] or '').strip()):>5}")
